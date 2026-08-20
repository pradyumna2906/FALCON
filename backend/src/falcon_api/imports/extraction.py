"""Bounded, non-evaluating CSV and XLSX statement extraction."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import PurePosixPath
from typing import Protocol, TypeAlias
from zipfile import BadZipFile, ZipFile

from falcon_api.core.errors import ApplicationError
from falcon_api.models.enums import ImportSourceType
from falcon_api.schemas.imports import StatementImportOptions
from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.cell.read_only import ReadOnlyCell


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_DATA_ROWS = 10_000
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_COMPRESSION_RATIO = 200
UPLOAD_CHUNK_BYTES = 64 * 1024

CSV_MEDIA_TYPES = frozenset({
    "application/csv",
    "text/csv",
    "text/plain",
})
XLSX_MEDIA_TYPES = frozenset({
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
})
_BLOCKED_ARCHIVE_PREFIXES = (
    "customxml/",
    "xl/activex/",
    "xl/embeddings/",
    "xl/externallinks/",
    "xl/oleobjects/",
)
_BLOCKED_ARCHIVE_NAMES = frozenset({"xl/vbaproject.bin"})
_REQUIRED_XLSX_PARTS = frozenset({"[content_types].xml", "xl/workbook.xml"})

CellScalar: TypeAlias = str | int | float | bool | date | datetime | time | None


class AsyncUpload(Protocol):
    """Minimal async upload interface used by FastAPI and unit tests."""

    async def read(self, size: int = -1) -> bytes:
        """Read at most size bytes."""


@dataclass(frozen=True, slots=True)
class ExtractedCell:
    """One inert cell value without executable formula content."""

    value: CellScalar
    formula_without_value: bool = False


@dataclass(frozen=True, slots=True)
class ExtractedRow:
    """One source row with its original one-based location."""

    row_number: int
    cells: tuple[ExtractedCell, ...]


@dataclass(frozen=True, slots=True)
class ExtractedStatement:
    """Bounded tabular source data ready for Phase 6.3 mapping."""

    header: ExtractedRow
    rows: tuple[ExtractedRow, ...]
    worksheet_name: str | None = None


async def read_bounded_upload(
    upload: AsyncUpload,
    *,
    max_bytes: int = MAX_UPLOAD_BYTES,
) -> bytes:
    """Read an async upload in chunks and fail after the first excess byte."""
    chunks: list[bytes] = []
    size = 0
    while True:
        chunk = await upload.read(min(UPLOAD_CHUNK_BYTES, max_bytes - size + 1))
        if not chunk:
            return b"".join(chunks)
        size += len(chunk)
        if size > max_bytes:
            raise _file_too_large()
        chunks.append(chunk)


def extract_statement(
    content: bytes,
    *,
    filename: str,
    content_type: str,
    options: StatementImportOptions,
    max_rows: int = MAX_DATA_ROWS,
) -> ExtractedStatement:
    """Validate source metadata and dispatch to the reviewed extractor."""
    media_type = content_type.partition(";")[0].strip().lower()
    if options.source_type == ImportSourceType.CSV:
        if not filename.lower().endswith(".csv") or media_type not in CSV_MEDIA_TYPES:
            raise _unsupported_file()
        return extract_csv(
            content,
            header_row=options.header_row,
            max_rows=max_rows,
        )
    if (
        not filename.lower().endswith(".xlsx")
        or media_type not in XLSX_MEDIA_TYPES
    ):
        raise _unsupported_file()
    return extract_xlsx(
        content,
        header_row=options.header_row,
        sheet_name=options.sheet_name,
        max_rows=max_rows,
    )


def extract_csv(
    content: bytes,
    *,
    header_row: int = 1,
    max_rows: int = MAX_DATA_ROWS,
) -> ExtractedStatement:
    """Decode and parse one bounded UTF-8 CSV with a reviewed delimiter."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise _file_too_large()
    if b"\x00" in content:
        raise _unsupported_file()
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise _unsupported_file() from None
    if not text.strip():
        raise _invalid_mapping()

    physical_lines = text.splitlines(keepends=True)
    if len(physical_lines) < header_row:
        raise _invalid_mapping()
    sample = "".join(physical_lines[header_row - 1 :])[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        parsed_rows = csv.reader(io.StringIO(text, newline=""), dialect, strict=True)
        return _collect_rows(parsed_rows, header_row=header_row, max_rows=max_rows)
    except (csv.Error, UnicodeError):
        raise _unsupported_file() from None


def extract_xlsx(
    content: bytes,
    *,
    header_row: int = 1,
    sheet_name: str | None = None,
    max_rows: int = MAX_DATA_ROWS,
) -> ExtractedStatement:
    """Read inert cached XLSX values after archive-security preflight."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise _file_too_large()
    _validate_xlsx_archive(content)

    try:
        formula_book = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        value_book = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (BadZipFile, KeyError, OSError, TypeError, ValueError):
        raise _unsupported_file() from None

    try:
        selected_name = sheet_name or formula_book.sheetnames[0]
        if selected_name not in formula_book.sheetnames:
            raise _invalid_mapping()
        formula_sheet = formula_book[selected_name]
        value_sheet = value_book[selected_name]
        rows = _iter_xlsx_rows(formula_sheet.iter_rows(), value_sheet.iter_rows())
        extracted = _collect_rows(rows, header_row=header_row, max_rows=max_rows)
        return ExtractedStatement(
            header=extracted.header,
            rows=extracted.rows,
            worksheet_name=selected_name,
        )
    finally:
        formula_book.close()
        value_book.close()


def _iter_xlsx_rows(
    formula_rows: Iterable[Sequence[Cell | ReadOnlyCell]],
    value_rows: Iterable[Sequence[Cell | ReadOnlyCell]],
) -> Iterator[tuple[ExtractedCell, ...]]:
    """Pair formula metadata with cached values without returning formulas."""
    for formula_row, value_row in zip(formula_rows, value_rows, strict=True):
        cells: list[ExtractedCell] = []
        for formula_cell, value_cell in zip(formula_row, value_row, strict=True):
            is_formula = formula_cell.data_type == "f"
            value = value_cell.value if is_formula else formula_cell.value
            cells.append(
                ExtractedCell(
                    value=value,
                    formula_without_value=is_formula and value is None,
                )
            )
        yield tuple(cells)


def _collect_rows(
    source_rows: Iterable[Sequence[str] | tuple[ExtractedCell, ...]],
    *,
    header_row: int,
    max_rows: int,
) -> ExtractedStatement:
    """Collect one header and a bounded number of physical data rows."""
    header: ExtractedRow | None = None
    rows: list[ExtractedRow] = []
    for row_number, source_row in enumerate(source_rows, start=1):
        cells = tuple(
            cell if isinstance(cell, ExtractedCell) else ExtractedCell(cell)
            for cell in source_row
        )
        row = ExtractedRow(row_number=row_number, cells=cells)
        if row_number < header_row:
            continue
        if row_number == header_row:
            header = row
            continue
        rows.append(row)
        if len(rows) > max_rows:
            raise _file_too_large()
    if header is None or not any(_cell_has_value(cell) for cell in header.cells):
        raise _invalid_mapping()
    return ExtractedStatement(header=header, rows=tuple(rows))


def _cell_has_value(cell: ExtractedCell) -> bool:
    value = cell.value
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _validate_xlsx_archive(content: bytes) -> None:
    """Reject encrypted, active, traversing, or explosively compressed XLSX."""
    if not content.startswith(b"PK"):
        raise _unsupported_file()
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_ARCHIVE_ENTRIES:
                raise _unsupported_file()
            names: set[str] = set()
            total_uncompressed = 0
            for entry in entries:
                normalized = entry.filename.replace("\\", "/").lower()
                path = PurePosixPath(normalized)
                if (
                    entry.flag_bits & 0x1
                    or path.is_absolute()
                    or ".." in path.parts
                    or normalized in names
                    or normalized in _BLOCKED_ARCHIVE_NAMES
                    or normalized.startswith(_BLOCKED_ARCHIVE_PREFIXES)
                ):
                    raise _unsupported_file()
                total_uncompressed += entry.file_size
                if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise _unsupported_file()
                has_impossible_size = (
                    entry.file_size > 0 and entry.compress_size == 0
                )
                has_unsafe_ratio = (
                    entry.compress_size > 0
                    and entry.file_size / entry.compress_size
                    > MAX_ARCHIVE_COMPRESSION_RATIO
                )
                if has_impossible_size or has_unsafe_ratio:
                    raise _unsupported_file()
                names.add(normalized)
            if not _REQUIRED_XLSX_PARTS.issubset(names):
                raise _unsupported_file()
    except (BadZipFile, OSError, ValueError):
        raise _unsupported_file() from None


def _file_too_large() -> ApplicationError:
    return ApplicationError(
        code="import_file_too_large",
        message="The import file exceeds the supported limit.",
        status_code=413,
    )


def _unsupported_file() -> ApplicationError:
    return ApplicationError(
        code="unsupported_import_file",
        message="The import file format is unsupported.",
        status_code=415,
    )


def _invalid_mapping() -> ApplicationError:
    return ApplicationError(
        code="invalid_import_mapping",
        message="The statement columns could not be mapped safely.",
        status_code=422,
    )
